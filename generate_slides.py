#!/usr/bin/env python3
"""
generate_slides.py
Generates a Product Development Process RACI PowerPoint presentation.
Run: pip install python-pptx && python generate_slides.py
Output: KICKR_Product_Development_RACI.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Colors ────────────────────────────────────────────────────────────────────
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF5, 0xF5, 0xF5)
LIGHT_ROW = RGBColor(0xEB, 0xEB, 0xEB)
DARK_BG   = RGBColor(0x12, 0x18, 0x2E)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY  = RGBColor(0xAA, 0xAA, 0xAA)
WAHOO_RED = RGBColor(0xE8, 0x20, 0x2F)

R_COL  = RGBColor(0xBF, 0x1F, 0x1F)
A_COL  = RGBColor(0xE0, 0x7C, 0x00)
C_COL  = RGBColor(0x1A, 0x65, 0xAD)
I_COL  = RGBColor(0x88, 0x88, 0x88)
RA_COL = RGBColor(0x7A, 0x18, 0x7A)

RACI_MAP = {"R": R_COL, "A": A_COL, "C": C_COL, "I": I_COL, "R/A": RA_COL}

PHASE_COLORS = [
    RGBColor(0x1B, 0x45, 0x7A),  # Discovery
    RGBColor(0x0D, 0x62, 0x8A),  # Concept
    RGBColor(0x0A, 0x74, 0x65),  # EVT Entry
    RGBColor(0x16, 0x7A, 0x3A),  # EVT
    RGBColor(0x5A, 0x72, 0x00),  # DVT Entry
    RGBColor(0x7A, 0x50, 0x00),  # DVT
    RGBColor(0x8B, 0x25, 0x00),  # PVT Handover
    RGBColor(0x8B, 0x00, 0x1A),  # PVT
    RGBColor(0x4A, 0x00, 0x4A),  # MP
]

# ── RACI Data ─────────────────────────────────────────────────────────────────
# Columns: SVP | Eng Dir | PM | Mech | Elec | FW | ID/UX | Prod Dev | Q/Mfg | Ops
PHASES = [
    {
        "name": "Discovery Phase",
        "color_idx": 0,
        "question": "Is this worth real concept work and deeper investment?",
        "activities": [
            ("Conduct Innovation Sessions",          ["I","I","R","C","C","C","C","A","",""]),
            ("Competitive Benchmarking",             ["I","I","I","R","R","R","I","A","I",""]),
            ("Technical Feasibility Assessments",    ["","A","I","R","R","R","","I","C",""]),
            ("Bench Top Concepts",                   ["","A","I","R","R","R","I","I","",""]),
            ("Platform Assessment",                  ["I","A","I","R","R","R","","C","",""]),
            ("System Mapping",                       ["","A","I","R","R","R","I","","",""]),
            ("Pre-BOM Inputs",                       ["","A","I","R","R","R","","I","","C"]),
            ("Past Problem History",                 ["I","I","I","I","I","I","I","I","R/A","I"]),
            ("Scope Regulations & Compliance",       ["","C","I","C","C","C","","R/A","I","I"]),
            ("Initial Make vs Buy Analysis",         ["I","C","I","C","C","C","","C","","R/A"]),
            ("IP Assessment",                        ["","C","I","C","C","C","C","I","","R/A"]),
            ("GATE: Discovery Exit Review",          ["A","R","R","R","R","R","R","R","R","R"]),
        ],
    },
    {
        "name": "Concept Phase",
        "color_idx": 1,
        "question": "Develop the idea to give the business clarity — should we invest further?",
        "activities": [
            ("ID Refinement",                        ["","C","I","C","C","","R/A","C","",""]),
            ("3D Print Evaluation",                  ["C","A","I","R","","","C","C","",""]),
            ("Initial CAD Assessment",               ["","A","I","R","C","","C","","",""]),
            ("Initial Failure Simulations",          ["","A","I","R","","","","","",""]),
            ("Initial DFM Review",                   ["I","R/A","I","","","","","C","C",""]),
            ("Schematic Outline",                    ["","C","I","I","R/A","C","","","",""]),
            ("Preliminary PCB Layout",               ["","C","I","I","R/A","C","","","",""]),
            ("Initial Firmware Development",         ["","I","I","","","R/A","","","",""]),
            ("Preliminary DFMEA",                    ["","A","I","R","R","R","","C","C",""]),
            ("Risk Register Established",            ["I","C","R/A","","","","C","C","C",""]),
            ("GATE: Concept Exit Review",            ["A","R","R","R","R","R","R","R","R","R"]),
        ],
    },
    {
        "name": "EVT Entry Expectations",
        "color_idx": 2,
        "question": "Do we have the required scope and is the design mature enough to begin EVT?",
        "activities": [
            ("Concept Agreement",                    ["C","R/A","I","C","C","C","R","C","",""]),
            ("BOM Review Complete",                  ["","R/A","I","C","C","C","","I","C","C"]),
            ("EVT Build Plan",                       ["","R/A","I","C","C","C","","","C","C"]),
            ("Concept CAD Frozen",                   ["","A","I","R","C","C","","","",""]),
            ("Schematic Frozen",                     ["","C","I","","R/A","C","","","",""]),
            ("Contract Manufacturing Confirmed",     ["I","C","","","","","","C","","R/A"]),
        ],
    },
    {
        "name": "EVT Phase",
        "color_idx": 3,
        "question": "Progressing design towards full validation and tooling release.",
        "activities": [
            ("Hardware Bring-Up",                    ["","A","I","","R","C","","","",""]),
            ("Firmware Development",                 ["","A","I","","C","R","","","",""]),
            ("Electrical Pre-Scan",                  ["","A","I","","R","","I","I","",""]),
            ("Assembly Verification",                ["","A","I","R","","","","","C",""]),
            ("DFM Finalized",                        ["","A","I","R","","","","","C","I"]),
            ("EoL Development",                      ["","A","I","","C","R","","I","I",""]),
            ("System Architecture & Feature Alloc.", ["","A","I","R","R","","","","",""]),
            ("DFMEA & Test Matrix Updated",          ["","R/A","I","C","C","C","","","C","C"]),
            ("Risk Register Updated",                ["I","C","R/A","","","","","","C","C"]),
            ("CM Reporting Cadence Agreement",       ["","C","R/A","","","","","","C","C"]),
            ("EVT Build",                            ["I","R/A","I","","","","","","C","C"]),
            ("GATE: EVT Exit Review",                ["A","R","R","R","R","R","R","R","R","R"]),
        ],
    },
    {
        "name": "DVT Entry Expectations",
        "color_idx": 4,
        "question": "Is the system architecture mature? Do we have a confirmed test plan?",
        "activities": [
            ("Design Status - Frozen",               ["I","A","R","C","C","C","C","C","C","C"]),
            ("Feature Functionality Complete",       ["","A","I","R","R","R","","C","",""]),
            ("Tooling Agreement",                    ["C","C","R","","","","","","C","A"]),
            ("BOM Review Complete",                  ["","R/A","I","C","C","C","","I","C","C"]),
            ("EVT Test Report Documented",           ["","A","I","R","R","R","","","C",""]),
            ("Certification Plan Documented",        ["I","C","R/A","","","","","I","C","C"]),
            ("Yield Report Reviewed",                ["I","I","I","","","","","","R/A","C"]),
        ],
    },
    {
        "name": "DVT Phase",
        "color_idx": 5,
        "question": "Progressing towards production level builds that can pass all validation and certification.",
        "activities": [
            ("Tooling Validation",                   ["I","C","I","","","","","","R/A","C"]),
            ("Reliability Testing",                  ["","C","I","","","","","","A","R"]),
            ("Process Flow Diagram Development",     ["","I","I","","","","","","R/A",""]),
            ("Establish Work Instructions",          ["","I","I","C","C","C","","","","R/A"]),
            ("Develop PFMEA",                        ["","R","I","C","C","C","","A","C",""]),
            ("Write Control Plan & Production Docs", ["","","I","","","","","","R/A","C"]),
            ("Determine Cosmetic Limit",             ["","C","I","","","","C","","R/A","C"]),
            ("DVT Build",                            ["I","R/A","I","","","","","","C","C"]),
            ("DVT Testing",                          ["I","R/A","I","","","","","","C","C"]),
            ("Regulatory Certification",             ["","A","R","","","","","I","",""]),
            ("Confirm Firmware Stability",           ["","I","I","","","R/A","","","",""]),
            ("Yield Analysis",                       ["","I","I","","","","","","R/A","C"]),
            ("Packaging Testing",                    ["","","I","","","","","","R/A",""]),
            ("GATE: DVT Exit Review",                ["A","R","R","R","R","R","R","R","R","R"]),
        ],
    },
    {
        "name": "DVT to PVT Handover",
        "color_idx": 6,
        "question": "Have we identified all engineering risks and provided support needed for manufacturing?",
        "activities": [
            ("Tooling Approval",                     ["","C","I","","","","","","R/A",""]),
            ("Regulatory Certifications Confirmed",  ["","A","R","","","","","I","",""]),
            ("Reliability Test Report Reviewed",     ["I","R/A","I","","","","","","C","C"]),
            ("Firmware Release Candidate Ready",     ["","I","I","","","R/A","","","C","C"]),
            ("Quality Plan Reviewed",                ["","","I","","","","","","R/A",""]),
            ("Run at Rate Validation",               ["","I","I","","","","","","R/A","C"]),
            ("Target Yield Met",                     ["","I","I","","","","","","R/A","C"]),
            ("GATE: PVT Handover Review",            ["A","R","R","R","R","R","R","R","R","R"]),
        ],
    },
    {
        "name": "PVT Phase",
        "color_idx": 7,
        "question": "Supporting activities: final BOM, certifications, PFMEA, work instructions.",
        "activities": [
            ("Final Component Cost Analysis",        ["","R/A","I","C","C","C","","I","C","C"]),
            ("PVT Build Complete",                   ["I","R/A","I","","","","","","C","C"]),
            ("Yield Root Cause Analysis",            ["","I","I","","","","","","R/A","C"]),
            ("Limit Sample Tuning",                  ["","C","I","","","","","","R/A","C"]),
            ("Maintain Firmware Stability",          ["","I","I","","","R/A","","","C",""]),
            ("Develop Service Manual / Repair Guide",["","R/A","I","","","","C","I","",""]),
            ("Technical Training",                   ["","R/A","I","","","","C","I","",""]),
            ("Complete After Action Review",         ["C","C","R/A","C","C","C","C","C","C","C"]),
        ],
    },
    {
        "name": "MP / Launch Exit Phase",
        "color_idx": 8,
        "question": "Are we moving from growth/launch into sustain/retire mode intentionally?",
        "activities": [
            ("Product Support / Sustainment Plan",   ["A","R","I","","","","","C","C","C"]),
        ],
    },
]

ROLE_HEADERS = ["SVP", "Eng Dir", "PM", "Mech", "Elec", "FW", "ID/UX", "Prod Dev", "Q/Mfg", "Ops"]

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_colored_box(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(
        1,  # MSO_AUTO_SHAPE_TYPE.RECTANGLE
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color  # match border to fill = invisible border
    return shape


def add_label(slide, text, left, top, width, height,
              size=12, bold=False, italic=False,
              color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def set_cell(cell, text, bg_color, text_color=WHITE,
             size=8, bold=False, align=PP_ALIGN.CENTER):
    cell.fill.solid()
    cell.fill.fore_color.rgb = bg_color
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = text_color


# ── Slide builders ────────────────────────────────────────────────────────────

def build_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_colored_box(slide, 0, 3.0, 13.33, 0.07, WAHOO_RED)
    add_label(slide, "Product Development Process",
              0.6, 1.5, 12.13, 1.2, size=36, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER)
    add_label(slide, "RACI Responsibility Matrix  |  Discovery through MP Launch",
              0.6, 2.8, 12.13, 0.55, size=15,
              color=MED_GRAY, align=PP_ALIGN.CENTER)

    legend = [("R  Responsible", R_COL), ("A  Accountable", A_COL),
              ("C  Consulted",   C_COL), ("I  Informed",    I_COL),
              ("R/A  Both",      RA_COL)]
    x = 0.92
    for label, col in legend:
        add_colored_box(slide, x, 5.8, 2.2, 0.55, col)
        add_label(slide, label, x, 5.8, 2.2, 0.55,
                  size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        x += 2.3


def build_overview_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)

    add_label(slide, "Development Phases at a Glance",
              0.3, 0.12, 12.73, 0.6, size=22, bold=True, color=WHITE)

    names = ["Discovery", "Concept", "EVT Entry", "EVT",
             "DVT Entry", "DVT", "PVT Handover", "PVT", "MP / Launch"]
    questions = [
        "Worth deeper\ninvestment?",    "Should we invest\nfurther?",
        "Design mature\nfor EVT?",      "Freeze arch,\npush to DVT?",
        "Arch mature?\nTest plan set?", "Pass validation\n& certs?",
        "Factory ready\nfor PVT?",      "Final BOM,\ncerts, training",
        "Sustain or\nretire mode?",
    ]
    bw = 13.33 / 9 - 0.07
    for i in range(9):
        x = 0.03 + i * (bw + 0.07)
        add_colored_box(slide, x, 0.85, bw, 1.0, PHASE_COLORS[i])
        add_label(slide, names[i], x, 0.88, bw, 0.45,
                  size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_label(slide, questions[i], x, 1.33, bw, 0.65,
                  size=7, color=RGBColor(0xDD, 0xDD, 0xDD), align=PP_ALIGN.CENTER)

    add_label(slide,
              "Gate reviews (Go / No-Go / Hold / Pivot) required at the end of each major phase",
              0.5, 2.1, 12.33, 0.4, size=10, italic=True,
              color=MED_GRAY, align=PP_ALIGN.CENTER)
    add_label(slide,
              "Roles:  SVP  |  Eng Directors  |  Project Management  |  Mechanical Eng  |  "
              "Electrical Eng  |  Firmware Eng  |  ID/UX  |  Product Development  |  "
              "Quality/Manufacturing  |  Operations",
              0.3, 2.65, 12.73, 0.45, size=9, color=MED_GRAY, align=PP_ALIGN.CENTER)

    legend = [("R  Responsible", R_COL), ("A  Accountable", A_COL),
              ("C  Consulted",   C_COL), ("I  Informed",    I_COL),
              ("R/A  Both",      RA_COL)]
    x = 1.42
    for label, col in legend:
        add_colored_box(slide, x, 3.3, 2.0, 0.42, col)
        add_label(slide, label, x, 3.3, 2.0, 0.42,
                  size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        x += 2.1


def build_phase_slide(prs, phase):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)

    pc = PHASE_COLORS[phase["color_idx"]]
    add_colored_box(slide, 0, 0, 13.33, 1.05, pc)
    add_label(slide, phase["name"], 0.2, 0.05, 9.5, 0.55,
              size=20, bold=True, color=WHITE)
    add_label(slide, phase["question"], 0.2, 0.6, 10.0, 0.38,
              size=9, italic=True, color=RGBColor(0xDD, 0xDD, 0xDD))

    lx = 9.7
    for lbl, col in [("R", R_COL), ("A", A_COL), ("C", C_COL),
                     ("I", I_COL), ("R/A", RA_COL)]:
        add_colored_box(slide, lx, 0.28, 0.6, 0.38, col)
        add_label(slide, lbl, lx, 0.28, 0.6, 0.38,
                  size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        lx += 0.68

    acts = phase["activities"]
    n_rows = len(acts) + 1
    n_cols = 11

    tbl = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(0.12), Inches(1.1),
        Inches(13.09), Inches(6.25)
    ).table

    tbl.columns[0].width = Inches(3.35)
    for ci in range(1, 11):
        tbl.columns[ci].width = Inches((13.09 - 3.35) / 10)

    # Header row
    set_cell(tbl.cell(0, 0), "Activity", pc, WHITE, size=8, bold=True,
             align=PP_ALIGN.LEFT)
    for ci, role in enumerate(ROLE_HEADERS):
        set_cell(tbl.cell(0, ci + 1), role, pc, WHITE, size=7, bold=True)

    # Activity rows
    for ri, (name, raci) in enumerate(acts):
        is_gate = name.startswith("GATE:")
        row_bg  = DARK_BG if is_gate else (OFF_WHITE if ri % 2 == 0 else LIGHT_ROW)
        txt_col = WHITE if is_gate else DARK_GRAY

        set_cell(tbl.cell(ri + 1, 0), name, row_bg, txt_col,
                 size=8, bold=is_gate, align=PP_ALIGN.LEFT)

        for ci, val in enumerate(raci):
            c = tbl.cell(ri + 1, ci + 1)
            if val in RACI_MAP:
                set_cell(c, val, RACI_MAP[val], WHITE, size=8, bold=True)
            else:
                set_cell(c, "", row_bg, txt_col, size=8)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    build_title_slide(prs)
    build_overview_slide(prs)
    for phase in PHASES:
        build_phase_slide(prs, phase)

    out = "KICKR_Product_Development_RACI.pptx"
    prs.save(out)
    print(f"Saved: {out}  ({len(PHASES) + 2} slides)")


if __name__ == "__main__":
    main()
